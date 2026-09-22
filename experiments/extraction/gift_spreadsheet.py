#!/usr/bin/env python3
"""Turn the full Arba Sicula extraction into the spreadsheet for Prof. Cipolla.

One row per aligned pair, with the issue, the printed Sicilian and English pages and
the LaBSE similarity, so every pair can be checked against the printed copy; a second
sheet counts the pairs per issue.

    python experiments/extraction/gift_spreadsheet.py \
        --tsv data/processed/as_full_gift/corpus.tsv \
        --out data/processed/as_full_gift/ArbaSicula_parallel_text.xlsx
"""
from __future__ import annotations
import argparse
import csv
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

REPO = Path(__file__).resolve().parents[2]


def issue_label(stem: str) -> str:
    """as04_05 -> '4-5' (double issue), as08_1 -> '8/1' (volume/number), as12 -> '12'."""
    parts = stem.removeprefix("as").split("_")
    if len(parts) == 1:
        return str(int(parts[0]))
    sep = "-" if len(parts[1]) == 2 else "/"
    return f"{int(parts[0])}{sep}{int(parts[1])}"


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--tsv", type=Path, default=REPO / "data/processed/as_full_gift/corpus.tsv")
    ap.add_argument("--out", type=Path,
                    default=REPO / "data/processed/as_full_gift/ArbaSicula_parallel_text.xlsx")
    args = ap.parse_args()

    with open(args.tsv, encoding="utf-8", newline="") as f:
        rows = list(csv.DictReader(f, delimiter="\t", quoting=csv.QUOTE_NONE))

    wb = Workbook()
    ws = wb.active
    ws.title = "Parallel text"
    ws.append(["Issue", "Sicilian page", "English page", "Sicilian", "English", "Similarity"])
    for r in rows:
        ws.append([issue_label(r["issue"]), r["scn_page"], r["en_page"],
                   r["sicilian"], r["english"], round(float(r["similarity"]), 3)])
    for col, width in zip("ABCDEF", (8, 10, 10, 70, 70, 10)):
        ws.column_dimensions[col].width = width
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=5):
        for cell in row:
            cell.alignment = wrap
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    per_issue = Counter(r["issue"] for r in rows)
    ss = wb.create_sheet("Pairs per issue")
    ss.append(["Issue", "Pairs"])
    for stem in sorted(per_issue):
        ss.append([issue_label(stem), per_issue[stem]])
    ss.append(["Total", len(rows)])
    for cell in ss[1]:
        cell.font = Font(bold=True)

    wb.save(args.out)
    print(f"{len(rows)} pairs from {len(per_issue)} issues -> {args.out}")


if __name__ == "__main__":
    main()
